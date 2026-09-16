"""Build a small corpus on disk.

Fixtures are generated, not committed. The PDF is assembled byte by byte so the tests
need no PDF writer and so the expected text is written down in one place.
"""

from __future__ import annotations

from pathlib import Path


def make_pdf(path: Path, pages: "list[list[str]]", title: str = "Fixture") -> None:
    """A minimal PDF with Helvetica text, one content stream per page."""
    objects: "list[bytes]" = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)  # object numbers are 1-based

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    catalog = add(b"")   # placeholder, filled once the pages object exists
    pages_obj = add(b"")
    info = add(
        "<< /Title ({}) /Producer (doc-duckdb fixture) >>".format(title).encode("latin-1")
    )

    kids = []
    for lines in pages:
        parts = [b"BT /F1 12 Tf 72 720 Td 14 TL"]
        for line in lines:
            escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            parts.append("({}) Tj T*".format(escaped).encode("latin-1"))
        parts.append(b"ET")
        stream = b"\n".join(parts)
        content = add(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream
            + b"\nendstream"
        )
        page = add(
            b"<< /Type /Page /Parent " + str(pages_obj).encode() + b" 0 R "
            b"/MediaBox [0 0 612 792] /Resources << /Font << /F1 "
            + str(font).encode() + b" 0 R >> >> /Contents "
            + str(content).encode() + b" 0 R >>"
        )
        kids.append(page)

    objects[catalog - 1] = (
        b"<< /Type /Catalog /Pages " + str(pages_obj).encode() + b" 0 R >>"
    )
    objects[pages_obj - 1] = (
        b"<< /Type /Pages /Count " + str(len(kids)).encode() + b" /Kids ["
        + b" ".join(str(k).encode() + b" 0 R" for k in kids) + b"] >>"
    )

    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"

    xref_at = len(out)
    out += "xref\n0 {}\n".format(len(objects) + 1).encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += "{:010d} 00000 n \n".format(offset).encode()
    out += (
        "trailer\n<< /Size {} /Root {} 0 R /Info {} 0 R >>\nstartxref\n{}\n%%EOF\n".format(
            len(objects) + 1, catalog, info, xref_at
        ).encode()
    )
    path.write_bytes(bytes(out))


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = "Docx fixture"
    document.core_properties.author = "Fixture author"
    document.add_heading("Heading one", level=1)
    document.add_paragraph("First body paragraph.")
    document.add_paragraph("   ")  # whitespace only: not a block
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "cell a"
    table.cell(0, 1).text = "cell b"
    table.cell(1, 0).text = "cell c"
    table.cell(1, 1).text = "cell d"
    document.add_paragraph("After the table.")
    document.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    presentation.core_properties.title = "Pptx fixture"
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(4572000), Emu(914400))
    box.text_frame.text = "Slide one title"
    box.text_frame.add_paragraph().text = "Slide one body"

    slide = presentation.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(4572000), Emu(914400))
    box.text_frame.text = "Slide two only"

    presentation.save(str(path))


def make_xlsx(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Numbers"
    sheet["A1"] = "label"
    sheet["B1"] = "value"
    sheet["A2"] = "alpha"
    sheet["B2"] = 12
    sheet["A3"] = "beta"
    sheet["B3"] = None
    second = workbook.create_sheet("Notes")
    second["A1"] = "a note"
    workbook.save(str(path))


def build(root: Path) -> Path:
    """Write the whole fixture corpus under root and return it."""
    root.mkdir(parents=True, exist_ok=True)
    (root / "nested").mkdir(exist_ok=True)

    make_pdf(root / "two_pages.pdf", [["Page one line one", "Page one line two"],
                                      ["Page two line one"]])
    make_pdf(root / "nested" / "one_page.pdf", [["Only line"]], title="Nested")
    make_docx(root / "sample.docx")
    make_pptx(root / "deck.pptx")
    make_xlsx(root / "book.xlsx")

    # failures that have to show up as rows, not as absences
    (root / "broken.pdf").write_bytes(b"this is not a pdf at all\n")
    (root / "broken.docx").write_bytes(b"not a zip, whatever the extension says")
    (root / "broken.xlsx").write_bytes(b"not a zip either")
    (root / "broken.pptx").write_bytes(b"nor this one")
    (root / "empty.pdf").write_bytes(b"")
    (root / "legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0 not parsed by this build")
    (root / "notes.txt").write_text("not a document format", encoding="utf-8")

    return root
